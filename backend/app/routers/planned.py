"""Planned data: plan lines (hosts to decommission per month), bulk import, audit trail.

Every mutation writes plan_audit rows in the same transaction as the change, with the
actor taken from the optional X-Changed-By header.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Dict, Optional

from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, UploadFile
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from .. import cache, config
from ..db import get_db
from ..models import PlanAudit, PlanLine, PlanMonthly, Portfolio

router = APIRouter(prefix="/api/planned", tags=["planned"])


class PlanLineIn(BaseModel):
    os: str
    portfolio: str
    application: str = ""
    owner: str = ""
    entry_date: Optional[str] = None
    months: Dict[str, int] = {}


# ------------------------------------------------------------------ helpers

def _mapped_portfolios(db: Session) -> dict[str, str]:
    """lower-case name -> canonical code, for every portfolio a plan may target."""
    return {p.code.lower(): p.code for p in db.query(Portfolio).all() if p.code != config.UNKNOWN}


def _parse_date(value: Optional[str]) -> date:
    if not value:
        return config.AS_OF_DATE
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        raise HTTPException(400, f"Bad entry date: {value}")


def _ident(line: PlanLine) -> dict:
    return dict(plan_line_id=line.id, os=line.os, portfolio=line.portfolio,
                application=line.application, owner=line.owner)


def _log(db: Session, action: str, ident: dict, field: str, old=None, new=None,
         source="ui", actor=None, detail=None) -> None:
    db.add(PlanAudit(action=action, field=field, old_value=None if old is None else str(old),
                     new_value=None if new is None else str(new), source=source,
                     actor=actor or None, detail=detail, ts=datetime.utcnow(), **ident))


def _validate(db: Session, body: PlanLineIn) -> PlanLineIn:
    if body.os not in config.TRACKED_OS:
        raise HTTPException(400, f"OS must be one of {config.TRACKED_OS}")
    canon = _mapped_portfolios(db).get(body.portfolio.strip().lower())
    if not canon:
        raise HTTPException(400, f"Unknown portfolio: {body.portfolio}")
    body.portfolio = canon
    valid = set(config.months())
    bad = [m for m in body.months if m not in valid]
    if bad:
        raise HTTPException(400, f"Months outside the reporting window: {', '.join(bad)}")
    if any(v < 0 for v in body.months.values()):
        raise HTTPException(400, "Planned counts can't be negative")
    return body


def _line_out(line: PlanLine, last: Optional[PlanAudit] = None) -> dict:
    return {
        "id": line.id, "os": line.os, "portfolio": line.portfolio,
        "application": line.application, "owner": line.owner,
        "entry_date": line.entry_date.isoformat(),
        "updated_at": line.updated_at.isoformat() if line.updated_at else None,
        "months": {pm.month: pm.planned_count for pm in line.months},
        "last_changed_by": last.actor if last else None,
        "last_changed_at": (last.ts if last else line.updated_at).isoformat() if (last or line.updated_at) else None,
    }


# ------------------------------------------------------------------ CRUD

@router.get("")
def list_lines(os: Optional[str] = None, db: Session = Depends(get_db)):
    q = db.query(PlanLine)
    if os:
        q = q.filter(PlanLine.os == os)
    lines = q.order_by(PlanLine.os, PlanLine.portfolio, PlanLine.owner, PlanLine.id).all()
    last: dict[int, PlanAudit] = {}
    for a in db.query(PlanAudit).order_by(PlanAudit.ts, PlanAudit.id).all():
        if a.plan_line_id is not None:
            last[a.plan_line_id] = a
    return [_line_out(l, last.get(l.id)) for l in lines]


@router.post("")
def create_line(body: PlanLineIn, db: Session = Depends(get_db),
                x_changed_by: Optional[str] = Header(None)):
    body = _validate(db, body)
    line = PlanLine(os=body.os, portfolio=body.portfolio, application=body.application.strip(),
                    owner=body.owner.strip(), entry_date=_parse_date(body.entry_date))
    line.months = [PlanMonthly(month=m, planned_count=v) for m, v in body.months.items() if v]
    db.add(line)
    db.flush()
    _log(db, "create", _ident(line), "line", new="created", actor=x_changed_by)
    for pm in line.months:
        _log(db, "create", _ident(line), "month:" + pm.month, new=pm.planned_count, actor=x_changed_by)
    db.commit()
    cache.invalidate()
    return _line_out(line)


@router.put("/{line_id}")
def update_line(line_id: int, body: PlanLineIn, db: Session = Depends(get_db),
                x_changed_by: Optional[str] = Header(None)):
    line = db.get(PlanLine, line_id)
    if not line:
        raise HTTPException(404, "Plan line not found")
    body = _validate(db, body)
    ident = _ident(line)  # log under the identity the line had before this edit
    changes = 0
    new_attrs = {"os": body.os, "portfolio": body.portfolio, "application": body.application.strip(),
                 "owner": body.owner.strip(), "entry_date": _parse_date(body.entry_date)}
    for field, new in new_attrs.items():
        old = getattr(line, field)
        if old != new:
            _log(db, "update", ident, field, old=old, new=new, actor=x_changed_by)
            setattr(line, field, new)
            changes += 1
    existing = {pm.month: pm for pm in line.months}
    for m in config.months():
        old = existing[m].planned_count if m in existing else 0
        new = body.months.get(m, 0)
        if old == new:
            continue
        _log(db, "update", ident, "month:" + m, old=old or None, new=new, actor=x_changed_by)
        changes += 1
        if m in existing:
            existing[m].planned_count = new
        else:
            line.months.append(PlanMonthly(month=m, planned_count=new))
    if changes:
        line.updated_at = datetime.utcnow()
        db.commit()
        cache.invalidate()
    return _line_out(line)


@router.delete("/{line_id}")
def delete_line(line_id: int, db: Session = Depends(get_db),
                x_changed_by: Optional[str] = Header(None)):
    line = db.get(PlanLine, line_id)
    if not line:
        raise HTTPException(404, "Plan line not found")
    total = sum(pm.planned_count for pm in line.months)
    _log(db, "delete", _ident(line), "line", old=f"{len(line.months)} month(s), total {total}",
         actor=x_changed_by)
    db.delete(line)
    db.commit()
    cache.invalidate()
    return {"deleted": line_id}


# ------------------------------------------------------------------ bulk import

_MONTH_NAMES = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def _parse_month_header(h: str) -> Optional[str]:
    """Accepts 2026-06, Jun-2026, Jun 2026, Jun '26 and similar."""
    s = h.strip().lower()
    m = re.fullmatch(r"(\d{4})-(\d{1,2})", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    m = re.fullmatch(r"([a-z]{3})[a-z]*[\s\-/']*'?(\d{2}|\d{4})", s)
    if m and m.group(1) in _MONTH_NAMES:
        y = int(m.group(2))
        y = y + 2000 if y < 100 else y
        return f"{y:04d}-{_MONTH_NAMES[m.group(1)]:02d}"
    return None


def _read_rows(raw: bytes, filename: str) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(it)]
        return [{header[i]: ("" if v is None else v) for i, v in enumerate(r) if i < len(header)}
                for r in it]
    text = raw.decode("utf-8-sig")
    dialect = csv.excel_tab if filename.lower().endswith(".tsv") else csv.excel
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


@router.post("/import")
async def import_plan(file: UploadFile = File(...), db: Session = Depends(get_db),
                      x_changed_by: Optional[str] = Header(None)):
    """Upsert plan lines from CSV/TSV/XLSX, keyed on (OS, Portfolio, Application, Owner)."""
    rows = _read_rows(await file.read(), file.filename or "upload.csv")
    if not rows:
        raise HTTPException(400, "The file has no rows")
    cols = {c.strip().lower(): c for c in rows[0].keys()}
    if "os" not in cols or "portfolio" not in cols:
        raise HTTPException(400, "Required columns: OS, Portfolio")
    window = set(config.months())
    month_cols = {c: _parse_month_header(c) for c in rows[0].keys()}
    month_cols = {c: m for c, m in month_cols.items() if m in window}

    mapped = _mapped_portfolios(db)
    known_owners = {o["username"] for o in _owners_cached(db)}
    existing = {(l.os, l.portfolio, l.application, l.owner): l for l in db.query(PlanLine).all()}
    months_cache: dict[int, dict[str, PlanMonthly]] = {}
    created = updated = skipped = 0
    seen: set = set()
    unmatched: dict[str, int] = {}
    unknown_owners: set[str] = set()
    detail = file.filename

    def get(r, name):
        c = cols.get(name)
        return str(r.get(c, "") if c else "").strip()

    for r in rows:
        os_name = get(r, "os")
        if os_name not in config.TRACKED_OS:
            skipped += 1
            continue
        pf_raw = get(r, "portfolio")
        pf = mapped.get(pf_raw.lower())
        if not pf:
            unmatched[pf_raw] = unmatched.get(pf_raw, 0) + 1
            skipped += 1
            continue
        app, owner = get(r, "application"), get(r, "owner")
        if owner and owner not in known_owners:
            unknown_owners.add(owner)
        key = (os_name, pf, app, owner)
        line = existing.get(key)
        if line is None:
            line = PlanLine(os=os_name, portfolio=pf, application=app, owner=owner,
                            entry_date=_parse_date(get(r, "entry date") or None))
            db.add(line)
            db.flush()
            existing[key] = line
            _log(db, "import", _ident(line), "line", new="created", source="import",
                 actor=x_changed_by, detail=detail)
            created += 1
        elif key not in seen:
            updated += 1
        seen.add(key)
        mc = months_cache.setdefault(line.id, {pm.month: pm for pm in line.months})
        for col, m in month_cols.items():
            val = str(r.get(col, "")).strip()
            if val == "":
                continue
            try:
                n = int(float(val))
            except ValueError:
                continue
            pm = mc.get(m)
            old = pm.planned_count if pm else None
            if old == n:
                continue
            if pm:
                pm.planned_count = n
            else:
                pm = PlanMonthly(plan_line_id=line.id, month=m, planned_count=n)
                db.add(pm)
                mc[m] = pm
            _log(db, "import", _ident(line), "month:" + m, old=old, new=n, source="import",
                 actor=x_changed_by, detail=detail)
        line.updated_at = datetime.utcnow()
    try:
        db.commit()
    except Exception as e:  # pragma: no cover - surfaced to the UI as a readable error
        db.rollback()
        raise HTTPException(400, f"Import failed: {e}")
    cache.invalidate()
    return {"created": created, "updated": updated, "skipped": skipped,
            "months_detected": sorted(set(month_cols.values())),
            "unmatched": [{"portfolio": k, "rows": v} for k, v in unmatched.items()],
            "unknown_owners": sorted(unknown_owners)}


def _owners_cached(db: Session) -> list[dict]:
    from .actuals import owners as owners_endpoint

    return owners_endpoint(db)


@router.get("/template")
def template():
    months = config.months()
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["OS", "Portfolio", "Application", "Owner", "Entry Date"] + months)
    w.writerow([config.TRACKED_OS[0], "SRE", "Example App", "username", config.AS_OF_DATE.isoformat()]
               + ["10" if i in (3, 4) else "" for i in range(len(months))])
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": 'attachment; filename="plan_template.csv"'})


@router.get("/unmatched")
def unmatched(db: Session = Depends(get_db)):
    """Plan lines whose portfolio doesn't match any known portfolio (won't join to actuals)."""
    mapped = set(_mapped_portfolios(db).values())
    out: dict = {}
    for l in db.query(PlanLine).all():
        if l.portfolio not in mapped:
            out[(l.os, l.portfolio)] = out.get((l.os, l.portfolio), 0) + 1
    return [{"os": k[0], "portfolio": k[1], "lines": v} for k, v in out.items()]


# ------------------------------------------------------------------ audit

def _audit_query(db: Session, action, os, portfolio, limit):
    q = db.query(PlanAudit)
    if action:
        q = q.filter(PlanAudit.action == action)
    if os:
        q = q.filter(PlanAudit.os == os)
    if portfolio:
        q = q.filter(PlanAudit.portfolio == portfolio)
    return q.order_by(PlanAudit.ts.desc(), PlanAudit.id.desc()).limit(limit).all()


_AUDIT_COLS = ["ts", "action", "os", "portfolio", "application", "owner", "field",
               "old_value", "new_value", "actor", "source", "detail", "plan_line_id"]


def _audit_out(a: PlanAudit) -> dict:
    d = {c: getattr(a, c) for c in _AUDIT_COLS}
    d["ts"] = a.ts.isoformat() if a.ts else None
    d["id"] = a.id
    return d


@router.get("/audit")
def audit(action: Optional[str] = None, os: Optional[str] = None, portfolio: Optional[str] = None,
          limit: int = Query(500, le=5000), db: Session = Depends(get_db)):
    return [_audit_out(a) for a in _audit_query(db, action, os, portfolio, limit)]


@router.get("/audit.csv")
def audit_csv(action: Optional[str] = None, os: Optional[str] = None, portfolio: Optional[str] = None,
              limit: int = Query(5000, le=5000), db: Session = Depends(get_db)):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(_AUDIT_COLS)
    for a in _audit_query(db, action, os, portfolio, limit):
        row = _audit_out(a)
        w.writerow([row[c] if row[c] is not None else "" for c in _AUDIT_COLS])
    buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": 'attachment; filename="plan_activity_log.csv"'})
